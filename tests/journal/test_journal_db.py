from __future__ import annotations

import threading
from datetime import date, datetime

import openpyxl
import pandas as pd
import pytest

from indexedge.core.clock import IST
from indexedge.core.types import (
    Fill,
    Instrument,
    InstrumentKind,
    NewsItem,
    Order,
    OrderRequest,
    OrderStatus,
    OrderType,
    Regime,
    Side,
    Signal,
    StrategySignal,
)
from indexedge.journal.db import SCHEMA_VERSION, Journal
from indexedge.journal.models import TradeRecord

FUT = Instrument(
    token="123", symbol="NIFTY26OCTFUT", name="NIFTY", exchange="NFO", kind=InstrumentKind.FUTURE, lot_size=75
)


def ist(d: int, h: int = 10, mi: int = 0) -> datetime:
    return datetime(2026, 10, d, h, mi, tzinfo=IST)


def make_signal(sid: str = "s1", strategy: str = "orb15", ts: datetime | None = None) -> Signal:
    raw = StrategySignal(
        strategy_id=strategy,
        underlying="NIFTY",
        ts=ts or ist(15),
        side=Side.BUY,
        entry=25000.0,
        stop=24950.0,
        targets=[25100.0],
        reasons=["Broke OR high"],
        features={"rsi": 61.2},
    )
    return Signal(
        id=sid,
        raw=raw,
        probability=0.55,
        probability_ci=(0.45, 0.65),
        sample_size=120,
        expected_rr=2.0,
        ev_per_trade=900.0,
        quantity=75,
        instrument=FUT,
        badges=["validated"],
        reasons=["ok"],
        context_reasons=["S&P +0.5%"],
        news_reasons=[],
        top_features=[("rsi", 0.12)],
        regime=Regime.TREND_UP,
        created_at=ts or ist(15),
    )


def test_schema_version_and_reopen(tmp_path):
    path = tmp_path / "j.duckdb"
    j = Journal(path)
    assert j.schema_version == SCHEMA_VERSION
    j.set_setting("x", {"a": 1})
    j.close()
    j2 = Journal(path)  # migrations are idempotent
    assert j2.schema_version == SCHEMA_VERSION
    assert j2.get_setting("x") == {"a": 1}
    j2.close()


def test_trade_record_validation(mk_trade):
    with pytest.raises(ValueError):
        mk_trade(mode="demo")
    with pytest.raises(ValueError):
        mk_trade(side="LONG")
    with pytest.raises(ValueError):
        mk_trade(qty=0)
    assert mk_trade(side=Side.SELL).side == "SELL"


def test_signals_roundtrip_and_status(journal):
    journal.record_signal(make_signal("s1"))
    journal.record_signal(make_signal("s2", strategy="vwap_pullback", ts=ist(16)))
    df = journal.signals()
    assert list(df["id"]) == ["s2", "s1"]  # newest first
    row = df[df["id"] == "s1"].iloc[0]
    assert row["targets"] == [25100.0]
    assert row["features"] == {"rsi": 61.2}
    assert row["instrument"] == "NIFTY26OCTFUT"
    assert row["ts"].tzinfo is not None and row["ts"].hour == 10
    journal.update_signal_status("s1", "skipped", "user skipped")
    assert journal.signals(status="skipped")["id"].tolist() == ["s1"]
    assert journal.signals(strategy="vwap_pullback")["id"].tolist() == ["s2"]
    assert journal.signals(start=date(2026, 10, 16))["id"].tolist() == ["s2"]
    assert len(journal.signals(limit=1)) == 1


def test_orders_upsert_and_fills(journal):
    req = OrderRequest(
        instrument=FUT, side=Side.BUY, quantity=75, order_type=OrderType.LIMIT, price=25001.0, tag="trd_1"
    )
    order = Order(request=req, broker_id="B1", created_at=ist(15))
    journal.record_order(order)
    order.status, order.filled_qty, order.avg_price = OrderStatus.FILLED, 75, 25000.5
    journal.record_order(order)
    df = journal.orders()
    assert len(df) == 1 and df.iloc[0]["status"] == "FILLED" and df.iloc[0]["filled_qty"] == 75
    journal.record_fill(Fill("B1", FUT, Side.BUY, 75, 25000.5, ist(15, 10, 1)))
    fills = journal.fills("B1")
    assert len(fills) == 1 and fills.iloc[0]["price"] == 25000.5


def test_open_close_trade_and_filters(journal, mk_trade):
    t1 = mk_trade()
    t2 = mk_trade(strategy_id="gap_fill", entry_ts=ist(16), side="SELL", entry_price=25000, stop=25040)
    journal.open_trade(t1)
    journal.open_trade(t2)
    journal.close_trade(t1.id, ist(15, 11), 25100.0, "target", pnl=7500.0, costs=60.0, mfe=110, mae=20)
    row = journal.trade(t1.id)
    assert row["status"] == "closed" and row["outcome"] == "target"
    assert row["r_multiple"] == pytest.approx(2.0)
    assert row["mfe"] == 110 and row["mae"] == 20
    journal.close_trade(t2.id, ist(16, 11), 25060.0, "stop", pnl=-4500.0)
    assert journal.trade(t2.id)["r_multiple"] == pytest.approx(-1.5)
    with pytest.raises(ValueError):
        journal.close_trade(t1.id, ist(15, 12), 1.0, "x", 0.0)
    with pytest.raises(KeyError):
        journal.close_trade("nope", ist(15, 12), 1.0, "x", 0.0)
    df = journal.trades()
    assert df["net_pnl"].tolist() == [7440.0, -4500.0]
    assert journal.trades(start=date(2026, 10, 15), end=date(2026, 10, 15))["id"].tolist() == [t1.id]
    assert journal.trades(strategy="gap_fill")["id"].tolist() == [t2.id]
    assert journal.trades(end=ist(15, 10, 0))["id"].tolist() == [t1.id]
    assert journal.trades(status="open").empty
    assert journal.trade("missing") is None


def test_today_stats(journal, mk_trade):
    day = date(2026, 10, 15)
    assert journal.today_stats(day).trades == 0
    pnls = [(1000.0, 10), (-500.0, 11), (-700.0, 12)]
    for i, (pnl, h) in enumerate(pnls):
        t = mk_trade(entry_ts=ist(15, h), qty=75 * (i + 1))
        journal.open_trade(t)
        journal.close_trade(t.id, ist(15, h, 30), 0.0 + 25000, "x", pnl=pnl, costs=50.0)
    bt = mk_trade(entry_ts=ist(15, 13), mode="backtest")
    journal.open_trade(bt)
    journal.close_trade(bt.id, ist(15, 13, 30), 24000.0, "stop", pnl=-99999.0)
    journal.open_trade(mk_trade(entry_ts=ist(15, 14), qty=150))
    s = journal.today_stats(day)
    assert s.trades == 4 and s.open_trades == 1
    assert s.realised_pnl == pytest.approx(1000 - 500 - 700 - 150)
    assert s.consecutive_losses == 2
    assert s.last_loss_ts == ist(15, 12, 30) and s.last_loss_qty == 225
    assert s.last_trade_qty == 150
    # a win resets the streak
    w = mk_trade(entry_ts=ist(15, 14, 10))
    journal.open_trade(w)
    journal.close_trade(w.id, ist(15, 14, 20), 25010.0, "target", pnl=500.0)
    assert journal.today_stats(day).consecutive_losses == 0
    assert journal.today_stats(day, include_backtest=True).trades == 6
    assert journal.today_stats(date(2026, 10, 16)).trades == 0


def test_notes_news_settings_audit_equity(journal):
    nid = journal.add_note("trd_1", "Entered late", ["fomo", "late"], "anxious", ts=ist(15))
    notes = journal.notes("trd_1")
    assert notes.iloc[0]["id"] == nid and notes.iloc[0]["tags"] == ["fomo", "late"]
    item = NewsItem(
        id="n1",
        ts=ist(15),
        source="rss",
        title="RBI holds repo rate",
        url="u",
        impact="high",
        entities=["RBI"],
    )
    journal.record_news(item)
    journal.record_news(item)  # dedupe by id
    journal.record_news(NewsItem(id="n2", ts=ist(16), source="rss", title="b", url="u2"))
    news = journal.news(limit=10)
    assert news["id"].tolist() == ["n2", "n1"] and news.iloc[1]["entities"] == ["RBI"]
    assert len(journal.news(limit=1)) == 1
    assert journal.get_setting("missing", 5) == 5
    journal.set_setting("k", [1, 2])
    journal.set_setting("k", [3])
    assert journal.get_setting("k") == [3]
    journal.audit_setting("risk.capital", 100000, 200000, ts=ist(14))
    aud = journal.settings_audit()
    assert aud.iloc[0]["old_value"] == 100000 and aud.iloc[0]["new_value"] == 200000
    journal.record_equity(date(2026, 10, 15), 201000.0, 1000.0, 2)
    journal.record_equity(date(2026, 10, 15), 202000.0, 2000.0, 3)
    eq = journal.equity_daily()
    assert len(eq) == 1 and eq.iloc[0]["equity"] == 202000.0


def test_export_csv_and_xlsx(journal, mk_trade, tmp_path):
    t = mk_trade(reasons=["a", "b"])
    journal.open_trade(t)
    journal.close_trade(t.id, ist(15, 11), 25050.0, "time", pnl=3750.0, costs=40.0)
    journal.record_signal(make_signal())
    journal.add_note(t.id, "fine", ["x"], "calm")
    csv = journal.export(tmp_path / "out" / "trades.csv", "csv")
    df = pd.read_csv(csv)
    assert df.iloc[0]["id"] == t.id and df.iloc[0]["reasons"] == '["a", "b"]'
    assert df.iloc[0]["entry_ts"].startswith("2026-10-15T10:00:00+05:30")
    xlsx = journal.export(tmp_path / "j.xlsx", "xlsx")
    wb = openpyxl.load_workbook(xlsx)
    assert {"trades", "signals", "notes", "orders", "fills"} <= set(wb.sheetnames)
    assert wb["trades"].max_row == 2 and wb["signals"].max_row == 2
    with pytest.raises(ValueError):
        journal.export(tmp_path / "x.pdf", "pdf")


def test_thread_safe_concurrent_writes(journal, mk_trade):
    def worker(k: int) -> None:
        for i in range(20):
            journal.open_trade(mk_trade(entry_ts=ist(15, 10, i)))
            journal.add_note(f"t{k}", "n")

    threads = [threading.Thread(target=worker, args=(k,)) for k in range(5)]
    for th in threads:
        th.start()
    for th in threads:
        th.join()
    assert len(journal.trades()) == 100
    assert len(journal.notes()) == 100


def test_trade_record_defaults():
    t = TradeRecord(None, "s", "NIFTY", "X", "BUY", 1, ist(15), 1.0, None)
    assert t.id.startswith("trd_") and t.status == "open" and t.targets == []
